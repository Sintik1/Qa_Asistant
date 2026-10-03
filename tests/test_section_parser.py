"""Hierarchical section parser — leaf units, steps, tables/code retention."""

from __future__ import annotations

import io

from docx import Document

from core.doc_reader import extract_text
from core.prompt_builder import build_leaf_user_prompt, iter_generation_prompts
from core.section_parser import ParseOptions, parse_requirements_document


SAMPLE = """
Вводный текст до разделов.

3. CRM
Описание системы CRM.

1. Доработать таблицу клиентов
2. Внести изменения в карточку
3. Изменить сроки действия акции в таблице

| Поле | Тип |
| Name | text |
| EndDate | date |

```sql
SELECT * FROM promo;
```

3.1 Каталог продуктов
Кратко про каталог.

1. Добавить фильтр по цене
2. Исправить сортировку

3.2 Карточка клиента

3.2.1 Редактирование полей
Текст подраздела.

1. Разрешить edit email

4. Billing
Раздел биллинга.

1. Проверить инвойс
"""


def test_leaf_sections_and_steps():
    parsed = parse_requirements_document(SAMPLE)
    numbers = [leaf.number for leaf in parsed.leaves]
    assert numbers == ["3", "3.1", "3.2.1", "4"]

    crm = next(leaf for leaf in parsed.leaves if leaf.number == "3")
    assert [i.text for i in crm.items] == [
        "Доработать таблицу клиентов",
        "Внести изменения в карточку",
        "Изменить сроки действия акции в таблице",
    ]
    assert "| Поле | Тип |" in crm.body
    assert "```sql" in crm.body
    assert "SELECT * FROM promo;" in crm.body

    catalog = next(leaf for leaf in parsed.leaves if leaf.number == "3.1")
    assert len(catalog.items) == 2
    assert "3.2 Карточка" in catalog.path or catalog.path.endswith("Каталог продуктов")

    deep = next(leaf for leaf in parsed.leaves if leaf.number == "3.2.1")
    assert "CRM" in deep.path
    assert "3.2 Карточка клиента" in deep.path
    assert deep.items[0].text.startswith("Разрешить")


def test_empty_container_not_emitted():
    parsed = parse_requirements_document(SAMPLE)
    # 3.2 has only a child, no own body/items → not emitted
    assert all(leaf.number != "3.2" for leaf in parsed.leaves)
    # 3 has own items + table before 3.1 → emitted as its own unit
    assert any(leaf.number == "3" for leaf in parsed.leaves)


def test_exclude_section_ids():
    parsed = parse_requirements_document(
        SAMPLE,
        options=ParseOptions(exclude_section_ids=frozenset({"4"})),
    )
    assert all(leaf.number != "4" for leaf in parsed.leaves)
    assert any(leaf.number == "3.1" for leaf in parsed.leaves)


def test_prompt_contains_path_items_and_body():
    parsed = parse_requirements_document(SAMPLE)
    crm = next(leaf for leaf in parsed.leaves if leaf.number == "3")
    prompt = build_leaf_user_prompt(crm)
    assert "Раздел (path):" in prompt
    assert "Доработать таблицу клиентов" in prompt
    assert "| EndDate | date |" in prompt
    assert "должны стать шагами" in prompt


def test_iter_prompts_one_per_leaf():
    parsed = parse_requirements_document(SAMPLE)
    prompts = iter_generation_prompts(parsed, raw_text_fallback=SAMPLE)
    assert len(prompts) == len(parsed.leaves)
    assert all(p[2] is not None for p in prompts)


def test_flat_fallback_without_sections():
    text = "Просто абзац требований без нумерации разделов про логин."
    parsed = parse_requirements_document(text)
    assert parsed.leaves == ()
    prompts = iter_generation_prompts(parsed, raw_text_fallback=text)
    assert len(prompts) == 1
    assert prompts[0][2] is None


def test_docx_keeps_table_inside_section_order():
    buf = io.BytesIO()
    doc = Document()
    doc.add_heading("3. CRM", level=1)
    doc.add_paragraph("1. Доработать таблицу")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Поле"
    table.cell(0, 1).text = "Тип"
    table.cell(1, 0).text = "Name"
    table.cell(1, 1).text = "text"
    doc.add_heading("3.1 Каталог", level=2)
    doc.add_paragraph("1. Добавить фильтр")
    doc.save(buf)

    extracted = extract_text("spec.docx", buf.getvalue())
    assert extracted.text.lstrip().startswith("#")
    parsed = parse_requirements_document(extracted.text)
    crm = next(leaf for leaf in parsed.leaves if leaf.number == "3")
    assert "Поле" in crm.body and "Name" in crm.body
    catalog = next(leaf for leaf in parsed.leaves if leaf.number == "3.1")
    assert "Поле" not in catalog.body
    assert catalog.items[0].text.startswith("Добавить")


def test_plain_numbering_without_heading_styles():
    """Normal-text outline ``3. / 3.1 / 3.2.1`` (no Word Heading / markdown)."""
    text = """
3. CRM
1. Доработать таблицу
2. Внести изменения

| Поле | Тип |
| A | B |

3.1 Каталог
1. Добавить фильтр

3.2 Карточка
3.2.1 Редактирование
1. Edit email

4. Billing
1. Проверить инвойс
"""
    parsed = parse_requirements_document(text)
    numbers = [leaf.number for leaf in parsed.leaves]
    assert numbers == ["3", "3.1", "3.2.1", "4"]
    crm = next(leaf for leaf in parsed.leaves if leaf.number == "3")
    assert len(crm.items) == 2
    assert "Поле" in crm.body


def test_mixed_heading_styles_and_numbering():
    """Heading styles and explicit numbers in one document."""
    text = """
## Основные требования
Общий контекст акции.

## СИСТЕМА 1
Вводный текст системы.

### 3.1 Каталог продуктов
1. Настроить EqType
| Код | Название |
| A | B |

### 3.2 Карточка
Текст карточки.

3.2.1 Редактирование полей
1. Разрешить edit email

## СИСТЕМА 2
Заказ недоступен.
"""
    parsed = parse_requirements_document(text)
    titles = {leaf.title: leaf for leaf in parsed.leaves}
    assert "Основные требования" in titles
    assert "СИСТЕМА 1" in titles
    assert "Каталог продуктов" in titles
    assert "Редактирование полей" in titles
    assert "СИСТЕМА 2" in titles
    catalog = titles["Каталог продуктов"]
    assert catalog.number == "3.1"
    assert "СИСТЕМА 1" in catalog.path
    assert "EqType" in catalog.items[0].text or "Настроить" in catalog.items[0].text
    assert "Код" in catalog.body
    deep = titles["Редактирование полей"]
    assert deep.number == "3.2.1"
    assert "Карточка" in deep.path


def test_markdown_heading_styles_like_trebovania():
    text = """
## Основные требования
Реализовать новую Акцию.
• согласование с клиентом;
• монтаж Оборудования;

| # | Тип | Цена |
| 1 | Роутер | 1000 |

## СИСТЕМА 1
Настроить новые продукты.

### Калькулятор
Настроить позиции в блоке WI-FI.
Для нового оборудования доступно отключение.

#### Информационные сообщения
• Добавить роутер в сообщения
• Настроить сообщения по акции

## СИСТЕМА 2
Заказ нового об-я по акции недоступно.
"""
    parsed = parse_requirements_document(text)
    titles = [leaf.title for leaf in parsed.leaves]
    assert "Основные требования" in titles
    assert "СИСТЕМА 1" in titles
    assert "Калькулятор" in titles
    assert "Информационные сообщения" in titles
    assert "СИСТЕМА 2" in titles

    main = next(leaf for leaf in parsed.leaves if leaf.title == "Основные требования")
    assert len(main.items) >= 2
    assert any("согласование" in i.text for i in main.items)
    assert "| Роутер |" in main.body or "Роутер" in main.body

    info = next(leaf for leaf in parsed.leaves if leaf.title == "Информационные сообщения")
    assert "СИСТЕМА 1" in info.path
    assert "Калькулятор" in info.path
    assert len(info.items) == 2
