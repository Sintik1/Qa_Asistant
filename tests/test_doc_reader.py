"""Unit tests for requirements document extraction."""

from __future__ import annotations

import io

import pytest
from docx import Document

from core.doc_reader import extract_text
from core.errors import ValidationError
from core.messages import ERROR_MESSAGES


def test_extract_markdown():
    data = "# Login\n\nUser can sign in with email.".encode("utf-8")
    result = extract_text("reqs.md", data)
    assert "sign in" in result.text
    assert result.extension == ".md"
    assert result.char_count >= 8


def test_extract_docx():
    buf = io.BytesIO()
    doc = Document()
    doc.add_heading("Requirements", level=1)
    doc.add_paragraph("User can reset password via email link.")
    doc.save(buf)
    result = extract_text("spec.docx", buf.getvalue())
    assert "reset password" in result.text
    assert result.extension == ".docx"


def test_invalid_format_exe():
    with pytest.raises(ValidationError) as exc:
        extract_text("virus.exe", b"MZ")
    assert exc.value.code == "INVALID_FORMAT"
    assert exc.value.message == ERROR_MESSAGES["INVALID_FORMAT"]


def test_empty_file():
    with pytest.raises(ValidationError) as exc:
        extract_text("empty.md", b"")
    assert exc.value.code == "EMPTY_FILE"


def test_no_requirements_too_short():
    with pytest.raises(ValidationError) as exc:
        extract_text("short.md", b"hi")
    assert exc.value.code == "NO_REQUIREMENTS"


def test_legacy_doc_ole_unsupported():
    # OLE compound signature (not a real Word file, but not ZIP/docx either)
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64
    with pytest.raises(ValidationError) as exc:
        extract_text("old.doc", ole)
    assert exc.value.code == "CORRUPT_FILE"


def test_corrupt_pdf():
    with pytest.raises(ValidationError) as exc:
        extract_text("bad.pdf", b"%PDF-1.4 not-a-real-pdf")
    assert exc.value.code == "CORRUPT_FILE"
